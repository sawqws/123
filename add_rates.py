"""Вставляет в лист «КАПИТАЛ» формулы Google Таблиц, которые сами подтягивают
курсы на дату из строки 32 (1-е число месяца).

USD/EUR/THB — официальный курс ЦБ РФ (cbr.ru, XML_daily), за 1 единицу валюты.
BTC/ETH — курс к USD из GOOGLEFINANCE × курс доллара ЦБ из той же колонки.
Будущие месяцы остаются пустыми, пока не наступит дата.

Запуск: python3 add_rates.py <вход.xlsx> <выход.xlsx>
"""
import sys

import openpyxl
from openpyxl.styles import PatternFill

SHEET = "КАПИТАЛ"
DATE_ROW = 32
RATE_COLS = ["C", "F", "I"]
CBR_ROWS = {33: "USD", 34: "EUR", 35: "THB"}
CRYPTO_ROWS = {36: "BTC", 37: "ETH"}
USD_ROW = 33
FORMULA_FILL = PatternFill("solid", fgColor="FFF2F2F2")
NOTE_CELL = "B39"
NOTE_ADD = (" Курсы (строки 33–37) подтягиваются сами на дату из строки 32: "
            "доллар, евро, бат — ЦБ РФ; BTC, ETH — GOOGLEFINANCE (к USD) × курс доллара ЦБ.")


def cbr_formula(col, code):
    url = f'"https://www.cbr.ru/scripts/XML_daily.asp?date_req="&TEXT({col}${DATE_ROW},"dd/mm/yyyy")'
    # Value у ЦБ всегда с 4 знаками после запятой: убираем запятую и делим на 10000,
    # чтобы не зависеть от десятичного разделителя в настройках таблицы.
    value = f"IMPORTXML({url},\"translate(//Valute[CharCode='{code}']/Value,',','')\")"
    nominal = f"IMPORTXML({url},\"//Valute[CharCode='{code}']/Nominal\")"
    return (f'=IF(OR({col}${DATE_ROW}="",{col}${DATE_ROW}>TODAY()),"",'
            f'IFERROR({value}/10000/{nominal},""))')


def crypto_formula(col, code):
    price = f'INDEX(GOOGLEFINANCE("CURRENCY:{code}USD","price",{col}${DATE_ROW}),2,2)'
    return (f'=IF(OR({col}${DATE_ROW}="",{col}${DATE_ROW}>TODAY(),N({col}{USD_ROW})=0),"",'
            f'IFERROR({price}*{col}{USD_ROW},""))')


def main(src, dst):
    wb = openpyxl.load_workbook(src)
    ws = wb[SHEET]
    for col in RATE_COLS:
        for row, code in CBR_ROWS.items():
            ws[f"{col}{row}"] = cbr_formula(col, code)
        for row, code in CRYPTO_ROWS.items():
            ws[f"{col}{row}"] = crypto_formula(col, code)
        for row in [*CBR_ROWS, *CRYPTO_ROWS]:
            ws[f"{col}{row}"].fill = FORMULA_FILL
    if NOTE_ADD not in (ws[NOTE_CELL].value or ""):
        ws[NOTE_CELL] = (ws[NOTE_CELL].value or "") + NOTE_ADD
    wb.save(dst)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
