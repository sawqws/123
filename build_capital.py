"""Собирает таблицу капитала из исходного xlsx:

1. У каждого месяца на всех листах появляется колонка «Комментарий»
   (как в отдельном файле «Доходы Расходы»).
2. На листе «КАПИТАЛ» курсы (строки 33–37) подтягиваются сами формулами
   Google Таблиц на 29-е число месяца колонки, 10:00 (в феврале — последний день).
   USD/EUR/THB — официальный курс ЦБ РФ (cbr.ru, XML_daily), за 1 единицу.
   BTC/ETH — курс к USD из GOOGLEFINANCE × курс доллара ЦБ из той же колонки.
   До наступления 29-го 10:00 ячейки пустые.

Запуск: python3 build_capital.py <вход.xlsx> <выход.xlsx>
"""
import copy
import re
import sys

import openpyxl
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.cell_range import CellRange

COMMENT_WIDTH = 26
FORMULA_FILL = PatternFill("solid", fgColor="FFF2F2F2")
INPUT_FILL = PatternFill("solid", fgColor="FFFFF8E1")

# Колонка «Сумма» каждого месяца в исходной таблице; комментарий встаёт сразу после неё.
# Есть ли у листа строка подзаголовков 2 («Кол-во / заметка», «Сумма, ₽»).
SHEETS = {
    "КАПИТАЛ": (["D", "G", "J"], True),
    "Плавающий капитал": (["D", "G", "J"], True),
    "Доходы Расходы": (["C", "E", "G"], False),
    "Долг": (["C", "E", "G"], False),
    "Авито": (["C", "E", "G"], False),
}

RATES_SHEET = "КАПИТАЛ"
DATE_ROW = 32
CBR_ROWS = {33: "USD", 34: "EUR", 35: "THB"}
CRYPTO_ROWS = {36: "BTC", 37: "ETH"}
USD_ROW = 33
RATE_DAY = 29
RATE_HOUR = 10

REF_RE = re.compile(r"(\$?)([A-Z]{1,3})(\$?)(\d+)")


def _split_sheet(ref):
    if "!" not in ref:
        return None, ref
    sheet, cells = ref.rsplit("!", 1)
    return sheet.strip("'"), cells


def _shift_formula(formula, own_sheet, target_sheet, idx):
    """Сдвигает вправо ссылки на колонки >= idx листа target_sheet."""
    tok = Tokenizer(formula)
    out = ["="]
    for t in tok.items:
        v = t.value
        if t.type == "OPERAND" and t.subtype == "RANGE":
            sheet, cells = _split_sheet(v)
            if (sheet or own_sheet) == target_sheet:
                def bump(m):
                    col = column_index_from_string(m.group(2))
                    if col >= idx:
                        col += 1
                    return f"{m.group(1)}{get_column_letter(col)}{m.group(3)}{m.group(4)}"
                cells = REF_RE.sub(bump, cells)
                v = (v[: len(v) - len(v.rsplit("!", 1)[1])] if sheet else "") + cells
        out.append(v)
    return "".join(out)


def _shift_range(rng, idx):
    r = CellRange(rng)
    if r.min_col >= idx:
        r.shift(col_shift=1)
    elif r.max_col >= idx:
        r.expand(right=1)
    return r.coord


def insert_col(wb, ws, idx):
    """Вставляет колонку idx на лист ws, сохраняя формулы, объединения, условное форматирование и ширины."""
    merges = [str(m) for m in ws.merged_cells.ranges]
    for m in merges:
        ws.unmerge_cells(m)
    widths = {k: v.width for k, v in ws.column_dimensions.items()}
    cf = ws.conditional_formatting

    ws.insert_cols(idx)

    for m in merges:
        ws.merge_cells(_shift_range(m, idx))
    for letter, w in widths.items():
        col = column_index_from_string(letter)
        ws.column_dimensions[get_column_letter(col + 1 if col >= idx else col)].width = w
    new_cf = ConditionalFormattingList()
    for rng, rules in cf._cf_rules.items():
        sqref = " ".join(_shift_range(str(r), idx) for r in rng.sqref.ranges)
        for rule in rules:
            new_cf.add(sqref, rule)
    ws.conditional_formatting = new_cf

    for sheet in wb:
        for row in sheet.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    c.value = _shift_formula(c.value, sheet.title, ws.title, idx)


def add_comment_columns(wb):
    for name, (sum_cols, has_subheader) in SHEETS.items():
        ws = wb[name]
        for letter in reversed(sum_cols):
            src = column_index_from_string(letter)
            dst = src + 1
            insert_col(wb, ws, dst)
            ws.column_dimensions[get_column_letter(dst)].width = COMMENT_WIDTH
            for row in range(1, ws.max_row + 1):
                s, d = ws.cell(row, src), ws.cell(row, dst)
                if s.has_style and (s.fill.fgColor.rgb not in (None, "00000000") or s.border.left.style):
                    d._style = copy.copy(s._style)
                    d.number_format = "General"
                    d.alignment = Alignment(wrap_text=True, vertical="center")
                    # Комментарий в строке статьи всегда для ввода, даже если сумма — формула.
                    if isinstance(ws.cell(row, 1).value, int):
                        d.fill = copy.copy(INPUT_FILL)
            header = ws.cell(1, src)
            sub = ws.cell(2, src) if has_subheader else None
            if has_subheader:
                # Дата месяца объединена над колонками месяца — растягиваем на комментарий.
                for m in list(ws.merged_cells.ranges):
                    if m.min_row == 1 and m.max_row == 1 and m.max_col == src:
                        ws.unmerge_cells(str(m))
                        ws.merge_cells(start_row=1, start_column=m.min_col, end_row=1, end_column=dst)
                ws.cell(2, dst).value = "Комментарий"
                ws.cell(2, dst)._style = copy.copy(sub._style)
            else:
                ws.cell(1, dst).value = "Комментарий"
                ws.cell(1, dst)._style = copy.copy(header._style)
                ws.cell(1, dst).number_format = "General"


def rate_cols(ws):
    """Колонки с курсами — те, где в строке дат курсов стоит ссылка на дату месяца."""
    return [c.column_letter for c in ws[DATE_ROW] if isinstance(c.value, str) and re.fullmatch(r"=\$?[A-Z]+\$?1", c.value or "")]


def date_formula(col):
    d = f"{col}$1"
    return (f"=DATE(YEAR({d}),MONTH({d}),MIN({RATE_DAY},DAY(EOMONTH({d},0))))"
            f"+TIME({RATE_HOUR},0,0)")


def cbr_formula(col, code):
    url = f'"https://www.cbr.ru/scripts/XML_daily.asp?date_req="&TEXT({col}${DATE_ROW},"dd/mm/yyyy")'
    # Value у ЦБ всегда с 4 знаками после запятой: убираем запятую и делим на 10000,
    # чтобы не зависеть от десятичного разделителя в настройках таблицы.
    value = f"IMPORTXML({url},\"translate(//Valute[CharCode='{code}']/Value,',','')\")"
    nominal = f"IMPORTXML({url},\"//Valute[CharCode='{code}']/Nominal\")"
    return (f'=IF(OR({col}${DATE_ROW}="",NOW()<{col}${DATE_ROW}),"",'
            f'IFERROR({value}/10000/{nominal},""))')


def crypto_formula(col, code):
    live = f'GOOGLEFINANCE("CURRENCY:{code}USD")'
    hist = f'INDEX(GOOGLEFINANCE("CURRENCY:{code}USD","price",{col}${DATE_ROW}),2,2)'
    # В сам день — текущая цена, со следующего дня — цена закрытия этого дня.
    return (f'=IF(OR({col}${DATE_ROW}="",NOW()<{col}${DATE_ROW},N({col}{USD_ROW})=0),"",'
            f'IFERROR(IF(INT({col}${DATE_ROW})=TODAY(),{live},{hist})*{col}{USD_ROW},""))')


def add_rates(wb):
    ws = wb[RATES_SHEET]
    for col in rate_cols(ws):
        ws[f"{col}{DATE_ROW}"] = date_formula(col)
        ws[f"{col}{DATE_ROW}"].number_format = "dd.mm.yy hh:mm"
        for row, code in CBR_ROWS.items():
            ws[f"{col}{row}"] = cbr_formula(col, code)
        for row, code in CRYPTO_ROWS.items():
            ws[f"{col}{row}"] = crypto_formula(col, code)
        for row in [*CBR_ROWS, *CRYPTO_ROWS]:
            ws[f"{col}{row}"].fill = FORMULA_FILL
    ws[f"B{DATE_ROW}"] = f"Курсы на {RATE_DAY}-е число, {RATE_HOUR}:00, ₽"


NOTES = {
    "КАПИТАЛ": "B39",
    "Плавающий капитал": "B14",
    "Доходы Расходы": "B43",
    "Долг": "B19",
    "Авито": "B19",
}
NOTE_COMMENT = " У каждого месяца своя колонка «Комментарий»."
NOTE_RATES = (" Курсы (строки 33–37) подтягиваются сами 29-го числа месяца в 10:00: "
              "доллар, евро, бат — ЦБ РФ; BTC, ETH — GOOGLEFINANCE (к USD) × курс доллара ЦБ.")


def rename_avito_comment(wb):
    ws = wb["Авито"]
    last = ws.cell(1, ws.max_column)
    if last.value == "Комментарий":
        last.value = "Общий комментарий"


def update_notes(wb):
    for name, cell in NOTES.items():
        ws = wb[name]
        text = ws[cell].value or ""
        if NOTE_COMMENT not in text:
            text += NOTE_COMMENT
        if name == RATES_SHEET and NOTE_RATES not in text:
            text += NOTE_RATES
        ws[cell] = text


def main(src, dst):
    wb = openpyxl.load_workbook(src)
    add_comment_columns(wb)
    rename_avito_comment(wb)
    add_rates(wb)
    update_notes(wb)
    wb.save(dst)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
